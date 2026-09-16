"""库存流水导出测试（阶段 5 CP2 交付物）。

覆盖 spec 验收项 10/11：
- 导出的 xlsx 能被 openpyxl.load_workbook 打开
- 行数 = 流水数 + 1（表头）
- 表头为中文
- 不走统一响应体（Content-Type 是 xlsx，响应体不是 JSON）
"""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook

from app.services.inventory_service import export_flows
from tests.test_api.conftest import login_as


@pytest.mark.integration
class TestExportService:
    """service 层导出（spec 验收 10）。"""

    def test_export_returns_valid_xlsx(self, db_session) -> None:
        """导出的字节流能被 openpyxl 打开，表头中文，行数 = 流水 + 1。"""
        content, filename = export_flows(db_session, store_id=1)
        assert filename.startswith("inventory_flows_")
        assert filename.endswith(".xlsx")
        # ASCII 文件名（无中文）
        assert filename.isascii()

        wb = load_workbook(io.BytesIO(content))
        ws = wb.active
        # 表头中文
        headers = [c.value for c in ws[1]]
        assert headers[0] == "发生时间"
        assert headers[1] == "商品名称"
        assert headers[2] == "流水类型"
        assert "操作人" in headers

        # 门店 1 有 20 条 INIT 流水 → 行数 = 20 + 1 表头
        assert ws.max_row == 21

        # 流水类型用中文 label
        first_type = ws.cell(row=2, column=3).value
        assert first_type == "期初库存"  # INIT 的中文

    def test_export_direction_chinese(self, db_session) -> None:
        """方向 1 → 入库。"""
        content, _ = export_flows(db_session, store_id=1)
        wb = load_workbook(io.BytesIO(content))
        ws = wb.active
        # INIT 流水 direction=1 → "入库"
        first_direction = ws.cell(row=2, column=4).value
        assert first_direction == "入库"


@pytest.mark.integration
class TestExportApi:
    """API 层导出（spec 验收 11：不走统一响应体）。"""

    def test_export_content_type_not_json(self, client) -> None:
        """响应 Content-Type 是 xlsx，响应体不是 JSON。"""

        data = login_as(client, "stocker01")
        headers = {"Authorization": f"Bearer {data['access_token']}"}
        r = client.get("/api/v1/inventory/flows/export", headers=headers)
        assert r.status_code == 200
        ct = r.headers.get("content-type", "")
        assert "spreadsheetml" in ct or "excel" in ct
        # 响应体是二进制 xlsx（PK 开头是 zip/xlsx 魔数），不是 JSON
        assert r.content[:2] == b"PK"
        # Content-Disposition 带 ASCII 文件名
        cd = r.headers.get("content-disposition", "")
        assert "inventory_flows_" in cd

    def test_export_requires_permission(self, client) -> None:
        """cashier01 无 inv:flow → 403。"""
        data = login_as(client, "cashier01")
        headers = {"Authorization": f"Bearer {data['access_token']}"}
        r = client.get("/api/v1/inventory/flows/export", headers=headers)
        assert r.status_code == 403
