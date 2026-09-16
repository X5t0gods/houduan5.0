"""库存业务编排（阶段 5 交付物）。

CP1：库存查询（/inventory）、流水台账（/inventory/flows）、三类预警（/inventory/warnings）
CP2：流水导出（/inventory/flows/export）

⚠️ 门店过滤（spec 5.1）：库存按门店隔离，取当前登录用户的 store_id；
   admin（data_scope=ALL）看全部门店（store_id=None）。
"""

from __future__ import annotations

import io
from datetime import datetime
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.error_codes import ErrorCode
from app.core.exceptions import BusinessError
from app.repositories import inventory_repository as inv_repo
from app.repositories import stock_flow_repository as flow_repo
from app.schemas.inventory import StockFlow, StockItem, WarningsResult

# 导出行数上限（spec 5.3：超出返回 2001，防内存爆）
EXPORT_MAX_ROWS = 10000

# flow_type → 中文 label（与前端 utils/dict.ts:83-93 FLOW_TYPE 一致，导出用）
FLOW_TYPE_LABELS = {
    "PURCHASE_IN": "采购入库",
    "SALE_OUT": "销售出库",
    "SALE_RETURN_IN": "销售退货入库",
    "PURCHASE_RETURN_OUT": "采购退货出库",
    "CHECK_ADJUST": "盘点调整",
    "LOSS_OUT": "报损出库",
    "TRANSFER_IN": "调拨入库",
    "TRANSFER_OUT": "调拨出库",
    "INIT": "期初库存",
}


def list_stock(
    session: Session,
    *,
    page: int = 1,
    page_size: int = 20,
    keyword: str | None = None,
    stock_status: str | None = None,
    store_id: int | None = None,
) -> tuple[list[StockItem], int]:
    """实时库存查询（spec 5.1）。"""
    rows, total = inv_repo.list_stock(
        session,
        page=page,
        page_size=page_size,
        keyword=keyword,
        stock_status=stock_status,
        store_id=store_id,
    )
    return [StockItem(**r) for r in rows], total


def list_flows(
    session: Session,
    *,
    page: int = 1,
    page_size: int = 20,
    keyword: str | None = None,
    flow_type: str | None = None,
    store_id: int | None = None,
    start_date: str | None = None,
    end_date: str | None = None,
) -> tuple[list[StockFlow], int]:
    """库存流水台账（spec 5.2）。"""
    rows, total = flow_repo.list_flows(
        session,
        page=page,
        page_size=page_size,
        keyword=keyword,
        flow_type=flow_type,
        store_id=store_id,
        start_date=start_date,
        end_date=end_date,
    )
    return [StockFlow(**r) for r in rows], total


def get_warnings(session: Session, store_id: int | None = None) -> WarningsResult:
    """三类预警（spec 5.4）。"""
    data = inv_repo.get_warnings(session, store_id)
    return WarningsResult(
        low_stock=[StockItem(**r) for r in data["low_stock"]],
        over_stock=[StockItem(**r) for r in data["over_stock"]],
        near_expiry=[StockItem(**r) for r in data["near_expiry"]],
    )


# ---------- 5.3 流水导出 ----------

# 导出表头（spec 5.3）
_EXPORT_HEADERS = [
    "发生时间", "商品名称", "流水类型", "方向", "变动数量",
    "变动前", "变动后", "单位成本", "变动金额", "来源单据", "操作人", "备注",
]


def export_flows(
    session: Session,
    *,
    keyword: str | None = None,
    flow_type: str | None = None,
    store_id: int | None = None,
    start_date: str | None = None,
    end_date: str | None = None,
) -> tuple[bytes, str]:
    """导出库存流水为 xlsx（spec 5.3）。

    ⚠️ 返回 (xlsx 字节流, 文件名)。文件名用 ASCII（避免浏览器下载乱码）。
    ⚠️ 超 EXPORT_MAX_ROWS 行 → 2001（防内存爆）。
    ⚠️ 流水类型用中文 label，方向 1→入库 / -1→出库。

    Returns:
        (bytes, filename)：filename 形如 inventory_flows_20260916.xlsx。
    """
    rows, total = flow_repo.list_flows_for_export(
        session,
        keyword=keyword,
        flow_type=flow_type,
        store_id=store_id,
        start_date=start_date,
        end_date=end_date,
        limit=EXPORT_MAX_ROWS + 1,
    )
    if total > EXPORT_MAX_ROWS:
        raise BusinessError(
            ErrorCode.REQUIRED_MISSING,
            f"导出数据超上限（{total} > {EXPORT_MAX_ROWS} 行），请缩小时间范围",
        )

    wb = Workbook()
    ws = wb.active
    ws.title = "库存流水"

    # 表头样式（蓝底白字加粗）
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4472C4")
    center = Alignment(horizontal="center", vertical="center")
    ws.append(_EXPORT_HEADERS)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    # 数据行
    for r in rows:
        created = r["created_at"]
        created_str = created.strftime("%Y-%m-%d %H:%M:%S") if created else ""
        direction_str = "入库" if r["direction"] == 1 else "出库"
        type_label = FLOW_TYPE_LABELS.get(r["flow_type"], r["flow_type"])
        ws.append([
            created_str,
            r["product_name"] or "",
            type_label,
            direction_str,
            float(r["quantity"]),
            float(r["before_qty"]),
            float(r["after_qty"]),
            float(r["unit_cost"]),
            float(r["amount"]),
            r["source_no"] or "",
            r["operator_name"] or "",
            r["remark"] or "",
        ])

    # 列宽
    widths = [20, 24, 14, 8, 12, 12, 12, 12, 12, 20, 12, 24]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    today = datetime.now(ZoneInfo(settings.TZ)).strftime("%Y%m%d")
    filename = f"inventory_flows_{today}.xlsx"
    return buf.getvalue(), filename
